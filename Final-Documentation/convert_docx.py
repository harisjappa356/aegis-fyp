import os
import random

try:
    import docx
except ImportError:
    print("python-docx is not installed. Please install it with 'pip install python-docx'")
    exit(1)

def convert_to_homoglyphs(text, replacement_chance=1.0):
    if not text:
        return text
    homoglyph_map = {
        'a': 'а', 'c': 'с', 'e': 'е', 'o': 'о', 'p': 'р', 'x': 'х', 'y': 'у',
        'A': 'А', 'B': 'Β', 'C': 'С', 'E': 'Ε', 'H': 'Н', 'M': 'Μ', 'O': 'Ο',
        'P': 'Ρ', 'T': 'Τ', 'X': 'Χ'
    }
    
    result = []
    for char in text:
        if char in homoglyph_map and random.random() <= replacement_chance:
            result.append(homoglyph_map[char])
        else:
            result.append(char)
            
    return ''.join(result)

filepath = r"D:\Final-Documentation\Final aegis-develop.docx"
output_filepath = r"D:\Final-Documentation\Final aegis-develop-homoglyphs.docx"

if not os.path.exists(filepath):
    print(f"File not found: {filepath}")
    exit(1)

print(f"Reading {filepath}...")
doc = docx.Document(filepath)

print("Converting paragraphs...")
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.text = convert_to_homoglyphs(run.text, 1.0)

print("Converting tables...")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.text = convert_to_homoglyphs(run.text, 1.0)

print(f"Saving to {output_filepath}...")
doc.save(output_filepath)
print("Conversion complete! Created: " + output_filepath)
